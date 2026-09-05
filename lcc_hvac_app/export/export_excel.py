from __future__ import annotations

import json
from io import BytesIO
from typing import Any

import pandas as pd
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Font

from lcc_hvac_app.engine.models import Project
from lcc_hvac_app.engine.scenario import compare_scenarios


FORMULA_REFERENCE = [
    {
        "Formula name": "Dust entering",
        "Formula": "Airflow x Dust concentration x Environment factor x Hours / 1000 / Qty",
        "Unit": "g/day/filter",
    },
    {
        "Formula name": "Dust captured",
        "Formula": "Dust entering x Mass efficiency",
        "Unit": "g/day/filter",
    },
    {
        "Formula name": "Estimated Mass Efficiency",
        "Formula": "PM1 fraction x ePM1 + PM1-2.5 fraction x ePM2.5 + PM2.5-10 fraction x ePM10 + Coarse fraction x ISO Coarse",
        "Unit": "%",
    },
    {
        "Formula name": "Face area",
        "Formula": "Width x Height / 1,000,000",
        "Unit": "m2/filter",
    },
    {
        "Formula name": "Airflow loading",
        "Formula": "Project airflow/filter / Rated airflow/filter x 100",
        "Unit": "%",
    },
    {
        "Formula name": "Face velocity",
        "Formula": "Airflow(m3/s) / (Face area x Qty/AHU)",
        "Unit": "m/s",
    },
    {
        "Formula name": "Media velocity",
        "Formula": "Airflow(m3/s) / (Media area x Qty/AHU)",
        "Unit": "m/s",
    },
    {
        "Formula name": "Filter life",
        "Formula": "DHC / Dust captured per day",
        "Unit": "days",
    },
    {
        "Formula name": "Replacement/year",
        "Formula": "Operating days/year / Filter life days",
        "Unit": "times/year",
    },
    {
        "Formula name": "Energy kWh",
        "Formula": "Airflow(m3/s) x Avg DP x Operating hours/year / (1000 x Fan efficiency) x AHU",
        "Unit": "kWh/year",
    },
    {
        "Formula name": "Energy cost",
        "Formula": "Energy kWh x Electricity price",
        "Unit": "VND/year",
    },
    {
        "Formula name": "TCO",
        "Formula": "Filter cost + Energy cost + Labor/disposal",
        "Unit": "VND/year",
    },
]


def _sheet_name(scenario_name: str) -> str:
    cleaned = scenario_name.replace("/", " ").replace("-", " ")
    parts = [part for part in cleaned.split() if part]
    return "_".join(parts)[:31]


def _project_summary_rows(project: Project, comparison: dict[str, Any]) -> list[dict[str, Any]]:
    info = project.project_info
    assumptions = project.assumptions
    summaries = comparison["summaries"]
    base = summaries[0] if summaries else {}
    best_name = comparison.get("best_option")
    best = next((row for row in summaries if row["scenario"] == best_name), base)
    return [
        {"Item": "Project name", "Value": info.project_name},
        {"Item": "Customer", "Value": info.customer},
        {"Item": "Location", "Value": info.location},
        {"Item": "Engineer", "Value": info.engineer},
        {"Item": "Date", "Value": info.date},
        {"Item": "Currency", "Value": info.currency},
        {"Item": "Airflow per AHU", "Value": assumptions.airflow_m3_h_per_ahu},
        {"Item": "Number of AHU", "Value": assumptions.number_of_ahu},
        {"Item": "Outdoor Environment", "Value": assumptions.outdoor_environment},
        {"Item": "Advanced dust override", "Value": assumptions.advanced_dust_override},
        {"Item": "Dust concentration (mg/m3)", "Value": assumptions.dust_concentration_mg_m3},
        {"Item": "Environment factor", "Value": assumptions.environment_factor},
        {"Item": "Best option", "Value": best_name},
        {"Item": "Base TCO/year", "Value": base.get("tco_year", 0)},
        {"Item": "Best TCO/year", "Value": best.get("tco_year", 0)},
        {"Item": "Saving/year", "Value": best.get("saving_vs_base", 0)},
        {"Item": "Saving %", "Value": best.get("saving_percent", 0)},
        {"Item": "3-year saving", "Value": best.get("saving_vs_base", 0) * 3},
        {"Item": "5-year saving", "Value": best.get("saving_vs_base", 0) * 5},
        {"Item": "CO2/year", "Value": best.get("co2_kg_year", 0)},
    ]


def _stage_input_rows(project: Project, scenario_name: str) -> list[dict[str, Any]]:
    scenario = next((item for item in project.scenarios if item.name == scenario_name), None)
    if scenario is None:
        return []
    rows = []
    for stage in scenario.stages:
        row = vars(stage).copy()
        row.setdefault("eurovent_iso_group", "ISO ePM1")
        row.setdefault("eurovent_mx_g", 200.0)
        row["eurovent_curve"] = json.dumps(
            getattr(stage, "eurovent_curve", []),
            ensure_ascii=False,
        )
        rows.append(row)
    return rows


def _write_chart_data(worksheet: Any, summaries_df: pd.DataFrame, stages_df: pd.DataFrame) -> None:
    summary_columns = [
        "scenario",
        "filter_cost_year",
        "energy_cost_year",
        "labor_disposal_cost_year",
        "tco_year",
        "energy_kwh_year",
        "co2_kg_year",
    ]
    stage_columns = ["scenario", "stage", "tco_year", "life_days"]

    worksheet["A1"] = "Scenario chart data"
    worksheet["A1"].font = Font(bold=True)
    for column_index, column_name in enumerate(summary_columns, start=1):
        cell = worksheet.cell(row=2, column=column_index, value=column_name)
        cell.font = Font(bold=True)
    for row_index, row in enumerate(
        summaries_df[summary_columns].itertuples(index=False), start=3
    ):
        for column_index, value in enumerate(row, start=1):
            worksheet.cell(row=row_index, column=column_index, value=value)

    stage_start_row = len(summaries_df) + 6
    worksheet.cell(row=stage_start_row, column=1, value="Stage chart data").font = Font(
        bold=True
    )
    for column_index, column_name in enumerate(stage_columns, start=1):
        cell = worksheet.cell(row=stage_start_row + 1, column=column_index, value=column_name)
        cell.font = Font(bold=True)
    for row_index, row in enumerate(
        stages_df[stage_columns].itertuples(index=False), start=stage_start_row + 2
    ):
        for column_index, value in enumerate(row, start=1):
            worksheet.cell(row=row_index, column=column_index, value=value)


def _add_charts(workbook: Any, summaries_df: pd.DataFrame, stages_df: pd.DataFrame) -> None:
    worksheet = workbook["Charts_Data"]
    scenario_count = len(summaries_df)
    stage_count = len(stages_df)
    if scenario_count == 0:
        return

    tco_chart = BarChart()
    tco_chart.type = "bar"
    tco_chart.style = 10
    tco_chart.title = "TCO Comparison"
    tco_chart.y_axis.title = "Scenario"
    tco_chart.x_axis.title = "Cost"
    data = Reference(worksheet, min_col=2, max_col=4, min_row=2, max_row=scenario_count + 2)
    categories = Reference(worksheet, min_col=1, min_row=3, max_row=scenario_count + 2)
    tco_chart.add_data(data, titles_from_data=True)
    tco_chart.set_categories(categories)
    tco_chart.shape = 4
    tco_chart.height = 8
    tco_chart.width = 16
    worksheet.add_chart(tco_chart, "I2")

    total_chart = BarChart()
    total_chart.title = "Total TCO by Scenario"
    total_chart.y_axis.title = "Cost"
    total_chart.x_axis.title = "Scenario"
    data = Reference(worksheet, min_col=5, min_row=2, max_row=scenario_count + 2)
    categories = Reference(worksheet, min_col=1, min_row=3, max_row=scenario_count + 2)
    total_chart.add_data(data, titles_from_data=True)
    total_chart.set_categories(categories)
    total_chart.height = 8
    total_chart.width = 16
    worksheet.add_chart(total_chart, "I18")

    co2_chart = BarChart()
    co2_chart.title = "CO2 Emission by Scenario"
    co2_chart.y_axis.title = "kgCO2/year"
    co2_chart.x_axis.title = "Scenario"
    data = Reference(worksheet, min_col=7, min_row=2, max_row=scenario_count + 2)
    categories = Reference(worksheet, min_col=1, min_row=3, max_row=scenario_count + 2)
    co2_chart.add_data(data, titles_from_data=True)
    co2_chart.set_categories(categories)
    co2_chart.height = 8
    co2_chart.width = 16
    worksheet.add_chart(co2_chart, "I34")

    if stage_count:
        stage_start_row = scenario_count + 6
        stage_header_row = stage_start_row + 1
        stage_first_row = stage_start_row + 2
        stage_last_row = stage_first_row + stage_count - 1

        stage_tco_chart = BarChart()
        stage_tco_chart.title = "TCO by Stage"
        stage_tco_chart.y_axis.title = "Cost"
        stage_tco_chart.x_axis.title = "Scenario / Stage"
        data = Reference(
            worksheet,
            min_col=3,
            min_row=stage_header_row,
            max_row=stage_last_row,
        )
        categories = Reference(worksheet, min_col=2, min_row=stage_first_row, max_row=stage_last_row)
        stage_tco_chart.add_data(data, titles_from_data=True)
        stage_tco_chart.set_categories(categories)
        stage_tco_chart.height = 8
        stage_tco_chart.width = 16
        worksheet.add_chart(stage_tco_chart, "Y2")

        life_chart = BarChart()
        life_chart.title = "Filter Life by Stage"
        life_chart.y_axis.title = "Days"
        life_chart.x_axis.title = "Scenario / Stage"
        data = Reference(
            worksheet,
            min_col=4,
            min_row=stage_header_row,
            max_row=stage_last_row,
        )
        categories = Reference(worksheet, min_col=2, min_row=stage_first_row, max_row=stage_last_row)
        life_chart.add_data(data, titles_from_data=True)
        life_chart.set_categories(categories)
        life_chart.height = 8
        life_chart.width = 16
        worksheet.add_chart(life_chart, "Y18")

    pie_chart = PieChart()
    pie_chart.title = "Base Cost Breakdown"
    data = Reference(worksheet, min_col=2, max_col=4, min_row=3, max_row=3)
    labels = Reference(worksheet, min_col=2, max_col=4, min_row=2, max_row=2)
    pie_chart.add_data(data, from_rows=True)
    pie_chart.set_categories(labels)
    pie_chart.dataLabels = DataLabelList()
    pie_chart.dataLabels.showPercent = True
    pie_chart.height = 8
    pie_chart.width = 12
    worksheet.add_chart(pie_chart, "Y34")


def build_excel_report(project: Project) -> bytes:
    comparison = compare_scenarios(project.scenarios, project.assumptions)
    summaries_df = pd.DataFrame(comparison["summaries"])
    stages_df = pd.DataFrame(comparison["stages"])
    assumptions_df = pd.DataFrame(
        [
            {"Parameter": key, "Value": value}
            for key, value in vars(project.assumptions).items()
        ]
    )
    filter_database_df = pd.DataFrame([vars(record) for record in project.filter_database])
    summary_df = pd.DataFrame(_project_summary_rows(project, comparison))
    formulas_df = pd.DataFrame(FORMULA_REFERENCE)

    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        summary_df.to_excel(writer, index=False, sheet_name="Project_Summary")
        assumptions_df.to_excel(writer, index=False, sheet_name="Assumptions")
        filter_database_df.to_excel(writer, index=False, sheet_name="Filter_Database")
        for scenario in project.scenarios:
            sheet = _sheet_name(scenario.name)
            pd.DataFrame(_stage_input_rows(project, scenario.name)).to_excel(
                writer, index=False, sheet_name=sheet
            )
        summaries_df.to_excel(writer, index=False, sheet_name="Scenario_Comparison")
        stages_df.to_excel(writer, index=False, sheet_name="Stage_Result")
        formulas_df.to_excel(writer, index=False, sheet_name="Formula_Reference")

        workbook = writer.book
        chart_sheet = workbook.create_sheet("Charts_Data")
        _write_chart_data(chart_sheet, summaries_df, stages_df)
        _add_charts(workbook, summaries_df, stages_df)
        for worksheet in workbook.worksheets:
            worksheet.freeze_panes = "A2"
            for column_cells in worksheet.columns:
                max_length = max(len(str(cell.value or "")) for cell in column_cells)
                worksheet.column_dimensions[column_cells[0].column_letter].width = min(
                    max(max_length + 2, 12), 46
                )
    output.seek(0)
    return output.getvalue()
