from io import BytesIO

from openpyxl import load_workbook

from lcc_hvac_app.engine.models import Assumptions, FilterStage, Project, Scenario
from lcc_hvac_app.export.export_excel import build_excel_report


def test_filter_life_excel_report_hides_dhc_only_parameters():
    project = Project(
        assumptions=Assumptions(calculation_method="Filter life-based"),
        scenarios=[
            Scenario(
                "Base / Current",
                [
                    FilterStage(
                        "Fine-filter",
                        qty_per_ahu=2,
                        dhc_g=0,
                        mass_efficiency=0,
                        avg_dp_pa=100,
                        price_vnd_filter=100000,
                        target_filter_life_days=180,
                    )
                ],
            )
        ],
    )

    workbook = load_workbook(BytesIO(build_excel_report(project)))
    assumption_items = [
        workbook["Assumptions"].cell(row=row, column=1).value
        for row in range(2, workbook["Assumptions"].max_row + 1)
    ]
    stage_headers = [cell.value for cell in workbook["Stage_Result"][1]]
    scenario_headers = [cell.value for cell in workbook["Base_Current"][1]]
    formula_names = [
        workbook["Formula_Reference"].cell(row=row, column=1).value
        for row in range(2, workbook["Formula_Reference"].max_row + 1)
    ]

    assert "outdoor_environment" not in assumption_items
    assert "dust_concentration_mg_m3" not in assumption_items
    assert "environment_factor" not in assumption_items
    assert "dhc_g" not in stage_headers
    assert "mass_efficiency" not in stage_headers
    assert "dhc_g" not in scenario_headers
    assert "mass_efficiency" not in scenario_headers
    assert "Dust entering" not in formula_names
    assert "Dust captured" not in formula_names
    assert "Estimated Mass Efficiency" not in formula_names
