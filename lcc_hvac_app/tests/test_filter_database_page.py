from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from lcc_hvac_app.ui.filter_database_page import (
    dataframe_to_excel_bytes,
    display_columns_for_calculation_method,
    filter_record_ids,
    _missing_filter_record_fields,
    normalize_uploaded_database,
    recalculate_media_area_dataframe,
    remove_filter_database_rows,
    visible_filter_database_dataframe,
)


def test_remove_filter_database_rows_removes_selected_ids():
    dataframe = pd.DataFrame(
        [
            {"Filter ID": "PF-001", "Supplier": "Air Filtech"},
            {"Filter ID": "FF-001", "Supplier": "Air Filtech"},
            {"Filter ID": "HEPA-001", "Supplier": "Air Filtech"},
        ]
    )

    remaining = remove_filter_database_rows(dataframe, ["PF-001", "HEPA-001"])

    assert remaining["Filter ID"].tolist() == ["FF-001"]


def test_filter_record_ids_skips_empty_values():
    dataframe = pd.DataFrame(
        [
            {"Filter ID": "PF-001"},
            {"Filter ID": ""},
            {"Filter ID": None},
            {"Filter ID": "FF-001"},
        ]
    )

    assert filter_record_ids(dataframe) == ["PF-001", "FF-001"]


def test_dataframe_to_excel_bytes_exports_filter_database_sheet():
    dataframe = pd.DataFrame(
        [
            {
                "Filter ID": "PF-001",
                "Supplier": "Air Filtech",
                "Stage": "Pre-filter",
                "Filter Class": "ISO Coarse 60%",
                "DHC/filter direct (g)": 450,
            }
        ]
    )

    workbook = load_workbook(filename=BytesIO(dataframe_to_excel_bytes(dataframe)))

    assert "Filter_Database" in workbook.sheetnames
    assert "_Lists" in workbook.sheetnames
    assert workbook["_Lists"].sheet_state == "hidden"
    worksheet = workbook["Filter_Database"]
    assert worksheet["A1"].value == "Filter ID"
    assert worksheet["A2"].value == "PF-001"
    assert "Filter Class" in [cell.value for cell in worksheet[1]]
    assert worksheet["A1"].fill.fgColor.rgb == "000F766E"
    assert worksheet["A2"].fill.fgColor.rgb == "00DCFCE7"
    assert worksheet.freeze_panes == "A2"
    assert len(worksheet.data_validations.dataValidation) >= 1
    headers = [cell.value for cell in worksheet[1]]
    media_area_column = headers.index("Media area/filter (m2)") + 1
    avg_dp_column = headers.index("Avg DP (Pa)") + 1
    media_area_cell = worksheet.cell(row=2, column=media_area_column)
    avg_dp_cell = worksheet.cell(row=2, column=avg_dp_column)
    assert media_area_cell.value.startswith("=IF(AND(")
    assert "ROUND(" in media_area_cell.value
    assert avg_dp_cell.value.startswith("=IF(AND(")
    assert "*1.1" in avg_dp_cell.value
    assert "ePM1 %" in [cell.value for cell in worksheet[1]]
    assert "Mass Eff. Source" in [cell.value for cell in worksheet[1]]
    assert any(
        validation.formula1.startswith("'_Lists'!$B$2:$B$")
        for validation in worksheet.data_validations.dataValidation
    )


def test_filter_life_database_view_hides_dhc_only_fields():
    dataframe = pd.DataFrame(
        [
            {
                "Filter ID": "FL-001",
                "Supplier": "Air Filtech",
                "Stage": "Fine-filter",
                "Filter Class": "E10",
                "Target filter life (days)": 180,
                "DHC/filter direct (g)": 600,
                "Mass Eff. %": 80,
                "ePM1 %": 50,
                "Price/filter (VND)": 500000,
            }
        ]
    )

    visible = visible_filter_database_dataframe(dataframe, "Filter life-based")

    assert visible.columns.tolist() == display_columns_for_calculation_method(
        "Filter life-based"
    )
    assert "Target filter life (days)" in visible.columns
    assert "DHC/filter direct (g)" not in visible.columns
    assert "Mass Eff. %" not in visible.columns
    assert "ePM1 %" not in visible.columns


def test_filter_life_excel_template_uses_filter_life_columns():
    dataframe = pd.DataFrame(
        [
            {
                "Filter ID": "FL-001",
                "Supplier": "Air Filtech",
                "Stage": "Fine-filter",
                "Filter Class": "E10",
                "Target filter life (days)": 180,
                "Avg DP (Pa)": 120,
                "Price/filter (VND)": 500000,
            }
        ]
    )

    workbook = load_workbook(
        filename=BytesIO(dataframe_to_excel_bytes(dataframe, "Filter life-based"))
    )
    headers = [cell.value for cell in workbook["Filter_Database"][1]]

    assert headers == display_columns_for_calculation_method("Filter life-based")
    assert "Target filter life (days)" in headers
    assert "DHC/filter direct (g)" not in headers
    assert "Mass Eff. %" not in headers
    assert "ePM1 %" not in headers


def test_dhc_excel_template_hides_filter_life_only_column():
    workbook = load_workbook(
        filename=BytesIO(dataframe_to_excel_bytes(pd.DataFrame(), "DHC-based"))
    )
    headers = [cell.value for cell in workbook["Filter_Database"][1]]

    assert headers == display_columns_for_calculation_method("DHC-based")
    assert "DHC/filter direct (g)" in headers
    assert "Mass Eff. %" in headers
    assert "Target filter life (days)" not in headers


def test_recalculate_media_area_dataframe_uses_width_and_height():
    dataframe = pd.DataFrame(
        [
            {
                "Filter ID": "PF-001",
                "Width (mm)": 592,
                "Height (mm)": 592,
                "Media area/filter (m2)": 0,
            }
        ]
    )

    normalized = recalculate_media_area_dataframe(dataframe)

    assert normalized.loc[0, "Media area/filter (m2)"] == 0.3505


def test_normalize_uploaded_database_imports_qty_per_ahu():
    dataframe = pd.DataFrame(
        [
            {
                "Filter ID": "PF-A",
                "Stage": "Pre-filter",
                "Filter Class": "ISO Coarse 70%",
                "Size": "592x287x46",
                "Qty/AHU": 20,
                "Rated airflow/filter (m3/h)": 1000,
                "Target filter life (days)": 180,
                "DHC to final DP (g)": 600,
                "Area (m2)": 1.2,
                "Eurovent Avg DP (Pa)": 85,
                "Mass Efficiency %": 70,
                "Price/filter (VND)": 250000,
            }
        ]
    )

    normalized = normalize_uploaded_database(dataframe)

    assert normalized.loc[0, "Filter ID"] == "PF-A"
    assert normalized.loc[0, "Filter Class"] == "ISO Coarse 70%"
    assert normalized.loc[0, "Qty/AHU"] == 20
    assert normalized.loc[0, "Rated airflow/filter (m3/h)"] == 1000
    assert normalized.loc[0, "Target filter life (days)"] == 180
    assert normalized.loc[0, "Width (mm)"] == 592
    assert normalized.loc[0, "Height (mm)"] == 287
    assert normalized.loc[0, "DHC/filter direct (g)"] == 600
    assert normalized.loc[0, "Media area/filter (m2)"] == 0.1699
    assert normalized.loc[0, "Avg DP (Pa)"] == 85
    assert normalized.loc[0, "Mass Eff. %"] == 70
    assert normalized.loc[0, "Price/filter (VND)"] == 250000


def test_normalize_uploaded_database_calculates_avg_dp_when_blank():
    dataframe = pd.DataFrame(
        [
            {
                "Filter ID": "PF-B",
                "Stage": "Pre-filter",
                "Initial DP (Pa)": 50,
                "Final DP (Pa)": 250,
            }
        ]
    )

    normalized = normalize_uploaded_database(dataframe)

    assert normalized.loc[0, "Avg DP (Pa)"] == 165


def test_normalize_uploaded_database_treats_nan_numeric_cells_as_blank():
    dataframe = pd.DataFrame(
        [
            {
                "Filter ID": "FL-NAN",
                "Stage": "Fine-filter",
                "Qty/AHU": float("nan"),
                "Target filter life (days)": float("nan"),
                "Width (mm)": float("nan"),
                "Height (mm)": 592,
                "Initial DP (Pa)": float("nan"),
                "Final DP (Pa)": 250,
                "Price/filter (VND)": float("nan"),
            }
        ]
    )

    normalized = normalize_uploaded_database(dataframe)

    assert normalized.loc[0, "Qty/AHU"] == 1
    assert normalized.loc[0, "Target filter life (days)"] == 0
    assert normalized.loc[0, "Width (mm)"] == 0
    assert normalized.loc[0, "Media area/filter (m2)"] == 0
    assert normalized.loc[0, "Avg DP (Pa)"] == 0
    assert normalized.loc[0, "Price/filter (VND)"] == 0


def test_normalize_uploaded_database_estimates_mass_efficiency_from_epm():
    dataframe = pd.DataFrame(
        [
            {
                "Filter ID": "FF-A",
                "Stage": "Fine-filter",
                "ePM10 %": 60,
                "ISO Coarse %": 90,
            }
        ]
    )

    normalized = normalize_uploaded_database(dataframe)

    assert normalized.loc[0, "Mass Eff. %"] == 70.5
    assert normalized.loc[0, "Mass Eff. Source"] == "Estimated from ePM/Coarse"


def test_normalize_uploaded_database_estimates_mass_efficiency_from_filter_class():
    dataframe = pd.DataFrame(
        [
            {
                "Filter ID": "FF-B",
                "Stage": "Fine-filter",
                "Filter Class": "ISO ePM1 50%",
            }
        ]
    )

    normalized = normalize_uploaded_database(dataframe)

    assert normalized.loc[0, "ePM1 %"] == 0
    assert normalized.loc[0, "Mass Eff. %"] == 50
    assert normalized.loc[0, "Mass Eff. Source"] == "Estimated from ePM/Coarse"


def test_normalize_uploaded_database_estimates_mass_efficiency_from_en1822_class():
    dataframe = pd.DataFrame(
        [
            {
                "Filter ID": "EPA-E11",
                "Stage": "EPA / Final-filter",
                "Filter Class": "E11",
            }
        ]
    )

    normalized = normalize_uploaded_database(dataframe)

    assert normalized.loc[0, "Mass Eff. %"] == 95
    assert normalized.loc[0, "Mass Eff. Source"] == "Estimated from EN1822 class"


def test_normalize_uploaded_database_still_accepts_legacy_iso_class_header():
    dataframe = pd.DataFrame(
        [
            {
                "Filter ID": "LEGACY",
                "Stage": "Fine-filter",
                "ISO Class": "ISO ePM10 60%",
            }
        ]
    )

    normalized = normalize_uploaded_database(dataframe)

    assert normalized.loc[0, "Filter Class"] == "ISO ePM10 60%"


def test_normalize_uploaded_database_detects_workbook_header_row():
    dataframe = pd.DataFrame(
        [
            ["FILTER DATABASE", None, None, None],
            ["Description", None, None, None],
            ["Filter ID", "Stage", "Qty/AHU", "DHC to final DP (g)"],
            ["PF-A", "Pre-filter", 20, 600],
        ]
    )

    normalized = normalize_uploaded_database(dataframe)

    assert normalized.loc[0, "Filter ID"] == "PF-A"
    assert normalized.loc[0, "Stage"] == "Pre-filter"
    assert normalized.loc[0, "Qty/AHU"] == 20
    assert normalized.loc[0, "DHC/filter direct (g)"] == 600


def test_filter_life_based_required_fields_do_not_require_dhc_or_efficiency(monkeypatch):
    monkeypatch.setattr(
        "lcc_hvac_app.ui.filter_database_page.current_calculation_method",
        lambda: "Filter life-based",
    )
    row = {
        "Filter ID": "FL-001",
        "Supplier": "Air Filtech",
        "Stage": "Fine-filter",
        "Filter model / description": "Life filter",
        "Filter Class": "E10",
        "Qty/AHU": 20,
        "Target filter life (days)": 180,
        "Avg DP (Pa)": 120,
        "Price/filter (VND)": 500000,
    }

    assert _missing_filter_record_fields(row) == []


def test_dhc_based_required_fields_require_dhc_and_efficiency(monkeypatch):
    monkeypatch.setattr(
        "lcc_hvac_app.ui.filter_database_page.current_calculation_method",
        lambda: "DHC-based",
    )
    row = {
        "Filter ID": "DHC-001",
        "Supplier": "Air Filtech",
        "Stage": "Fine-filter",
        "Filter model / description": "DHC filter",
        "Filter Class": "",
        "Qty/AHU": 20,
        "Target filter life (days)": 180,
        "Avg DP (Pa)": 120,
        "Price/filter (VND)": 500000,
    }

    missing = _missing_filter_record_fields(row)

    assert "DHC/filter direct (g)" in missing
    assert "Mass Eff. % or Filter Class/ePM/Coarse" in missing
