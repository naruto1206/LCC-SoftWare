from __future__ import annotations

import re
from io import BytesIO
from typing import Any

import pandas as pd
import streamlit as st
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from lcc_hvac_app.engine.efficiency import (
    DEFAULT_OUTDOOR_ENVIRONMENT,
    effective_mass_efficiency,
    environment_profile,
    percent_for_display,
)
from lcc_hvac_app.engine.models import FilterDatabaseRecord


DATABASE_COLUMNS = [
    "filter_id",
    "supplier",
    "stage",
    "model",
    "filter_class",
    "qty_per_ahu",
    "rated_airflow_m3_h_filter",
    "width_mm",
    "height_mm",
    "media_area_m2",
    "dhc_g",
    "initial_dp_pa",
    "avg_dp_pa",
    "final_dp_pa",
    "epm1_percent",
    "epm25_percent",
    "epm10_percent",
    "coarse_percent",
    "mass_efficiency",
    "mass_efficiency_source",
    "price_vnd_filter",
    "notes",
]

DISPLAY_COLUMNS = {
    "filter_id": "Filter ID",
    "supplier": "Supplier",
    "stage": "Stage",
    "model": "Filter model / description",
    "filter_class": "ISO Class",
    "qty_per_ahu": "Qty/AHU",
    "rated_airflow_m3_h_filter": "Rated airflow/filter (m3/h)",
    "width_mm": "Width (mm)",
    "height_mm": "Height (mm)",
    "media_area_m2": "Media area/filter (m2)",
    "dhc_g": "DHC/filter direct (g)",
    "initial_dp_pa": "Initial DP (Pa)",
    "avg_dp_pa": "Avg DP (Pa)",
    "final_dp_pa": "Final DP (Pa)",
    "epm1_percent": "ePM1 %",
    "epm25_percent": "ePM2.5 %",
    "epm10_percent": "ePM10 %",
    "coarse_percent": "ISO Coarse %",
    "mass_efficiency": "Mass Eff. %",
    "mass_efficiency_source": "Mass Eff. Source",
    "price_vnd_filter": "Price/filter (VND)",
    "notes": "Notes",
}

STAGE_OPTIONS = [
    "",
    "Pre-filter",
    "Fine-filter",
    "EPA / Final-filter",
    "HEPA",
    "ULPA",
]

ISO_CLASS_OPTIONS = (
    [""]
    + [f"ISO Coarse {value}%" for value in range(10, 101, 5)]
    + [f"ISO ePM10 {value}%" for value in range(50, 101, 5)]
    + [f"ISO ePM2.5 {value}%" for value in range(50, 101, 5)]
    + [f"ISO ePM1 {value}%" for value in range(50, 101, 5)]
)

COLUMN_ALIASES = {
    "filter_id": ["filter id", "id", "code", "filter code", "product code"],
    "supplier": ["supplier", "supplier/brand", "brand", "manufacturer", "maker"],
    "stage": ["stage", "filter stage", "level"],
    "model": ["model", "filter model", "filter model / description", "description", "product", "item"],
    "filter_class": ["iso class", "class", "grade", "filter class", "en class"],
    "qty_per_ahu": ["qty/ahu", "qty ahu", "quantity/ahu", "quantity per ahu", "qty per ahu"],
    "rated_airflow_m3_h_filter": [
        "rated airflow/filter",
        "rated airflow/filter (m3/h)",
        "rated airflow per filter",
        "rated airflow per filter (m3/h)",
        "nominal airflow",
        "nominal airflow/filter",
        "airflow/filter",
        "airflow per filter",
        "air flow",
        "airflow",
    ],
    "width_mm": ["width", "width mm", "width (mm)", "w", "w mm"],
    "height_mm": ["height", "height mm", "height (mm)", "h", "h mm"],
    "media_area_m2": [
        "media area",
        "media area/filter",
        "media area/filter m2",
        "media area/filter m²",
        "media area/filter (m2)",
        "media area (m2)",
        "area",
        "area (m2)",
    ],
    "dhc_g": [
        "dhc",
        "dhc (g)",
        "dhc/filter direct g",
        "dhc/filter direct (g)",
        "dhc to final dp (g)",
        "dust holding capacity",
        "dust holding capacity (g)",
    ],
    "initial_dp_pa": ["initial dp", "initial dp (pa)", "initial pressure drop"],
    "avg_dp_pa": [
        "avg dp",
        "avg dp (pa)",
        "average dp",
        "average pressure drop",
        "eurovent avg dp (pa)",
    ],
    "final_dp_pa": ["final dp", "final dp (pa)", "final pressure drop"],
    "epm1_percent": ["epm1", "epm1 %", "epm1%", "iso epm1", "iso epm1 %"],
    "epm25_percent": [
        "epm2.5",
        "epm2.5 %",
        "epm2.5%",
        "epm25",
        "epm25 %",
        "iso epm2.5",
    ],
    "epm10_percent": ["epm10", "epm10 %", "epm10%", "iso epm10", "iso epm10 %"],
    "coarse_percent": [
        "coarse",
        "coarse %",
        "coarse%",
        "iso coarse",
        "iso coarse %",
        "arrestance",
        "arrestance %",
    ],
    "mass_efficiency": [
        "mass efficiency",
        "mass efficiency %",
        "mass eff.",
        "mass eff. %",
        "mass eff %",
        "efficiency",
        "eff",
    ],
    "mass_efficiency_source": ["mass eff. source", "mass efficiency source", "efficiency source"],
    "price_vnd_filter": [
        "price/filter",
        "price/filter vnd",
        "price/filter (vnd)",
        "price",
        "unit price",
        "filter price",
    ],
    "notes": ["notes", "note", "remark", "remarks"],
}


def _choice_options(defaults: list[str], dataframe: pd.DataFrame, display_column: str) -> list[str]:
    options = list(defaults)
    if display_column in dataframe.columns:
        for value in dataframe[display_column].dropna():
            text = str(value).strip()
            if text and text not in options:
                options.append(text)
    return options


def records_to_dataframe(records: list[FilterDatabaseRecord]) -> pd.DataFrame:
    data = [{column: getattr(record, column) for column in DATABASE_COLUMNS} for record in records]
    return pd.DataFrame(data, columns=DATABASE_COLUMNS).rename(columns=DISPLAY_COLUMNS)


def empty_display_row() -> dict[str, Any]:
    return {
        column: 0.0
        if column
        in {
            "DHC (g)",
            "DHC/filter direct (g)",
            "Qty/AHU",
            "Rated airflow/filter (m3/h)",
            "Width (mm)",
            "Height (mm)",
            "Media area/filter (m2)",
            "Initial DP (Pa)",
            "Avg DP (Pa)",
            "Final DP (Pa)",
            "ePM1 %",
            "ePM2.5 %",
            "ePM10 %",
            "ISO Coarse %",
            "Mass Efficiency",
            "Mass Eff. %",
            "Price/filter",
            "Price/filter (VND)",
        }
        else ""
        for column in DISPLAY_COLUMNS.values()
    }


def calculate_media_area_m2(width_mm: Any, height_mm: Any) -> float:
    width = _to_float(width_mm)
    height = _to_float(height_mm)
    if width <= 0 or height <= 0:
        return 0.0
    return round(width * height / 1_000_000, 4)


def parse_iso_class_efficiencies(iso_class: Any) -> dict[str, float]:
    text = str(iso_class or "").strip()
    match = re.search(r"(\d+(?:\.\d+)?)\s*%", text)
    if match is None:
        return {}
    value = _to_float(match.group(1))
    lower_text = text.lower().replace(",", ".")
    if "epm2.5" in lower_text or "epm25" in lower_text:
        return {"ePM2.5 %": value}
    if "epm10" in lower_text:
        return {"ePM10 %": value}
    if "epm1" in lower_text:
        return {"ePM1 %": value}
    if "coarse" in lower_text:
        return {"ISO Coarse %": value}
    return {}


def current_outdoor_environment() -> str:
    project = st.session_state.get("project")
    assumptions = getattr(project, "assumptions", None)
    return getattr(assumptions, "outdoor_environment", DEFAULT_OUTDOOR_ENVIRONMENT)


def calculate_effective_mass_efficiency_for_row(row: pd.Series | dict[str, Any]) -> tuple[float, str]:
    source = str(row.get("Mass Eff. Source", "") or "")
    direct = 0.0 if source.startswith("Estimated") else row.get("Mass Eff. %", 0.0)
    parsed_iso = parse_iso_class_efficiencies(row.get("ISO Class", ""))
    epm1 = row.get("ePM1 %", 0.0) or parsed_iso.get("ePM1 %", 0.0)
    epm25 = row.get("ePM2.5 %", 0.0) or parsed_iso.get("ePM2.5 %", 0.0)
    epm10 = row.get("ePM10 %", 0.0) or parsed_iso.get("ePM10 %", 0.0)
    coarse = row.get("ISO Coarse %", 0.0) or parsed_iso.get("ISO Coarse %", 0.0)
    return effective_mass_efficiency(
        direct,
        epm1,
        epm25,
        epm10,
        coarse,
        current_outdoor_environment(),
    )


def recalculate_mass_efficiency_dataframe(dataframe: pd.DataFrame) -> pd.DataFrame:
    normalized = dataframe.copy()
    for column in ["ePM1 %", "ePM2.5 %", "ePM10 %", "ISO Coarse %", "Mass Eff. %"]:
        if column in normalized.columns:
            normalized[column] = normalized[column].apply(_to_float)
    if "Mass Eff. Source" not in normalized.columns:
        normalized["Mass Eff. Source"] = ""
    if "Mass Eff. %" not in normalized.columns:
        return normalized
    calculated_values: list[float] = []
    sources: list[str] = []
    for _, row in normalized.iterrows():
        efficiency, source = calculate_effective_mass_efficiency_for_row(row)
        calculated_values.append(round(efficiency * 100.0, 4) if source != "Missing" else 0.0)
        sources.append(source)
    normalized["Mass Eff. %"] = calculated_values
    normalized["Mass Eff. Source"] = sources
    return normalized


def normalize_display_dataframe(dataframe: pd.DataFrame | None) -> pd.DataFrame:
    if dataframe is None:
        return pd.DataFrame(columns=list(DISPLAY_COLUMNS.values()))
    normalized = dataframe.copy()
    if "Model" in normalized.columns and "Filter model / description" not in normalized.columns:
        normalized = normalized.rename(columns={"Model": "Filter model / description"})
    if "Class" in normalized.columns and "ISO Class" not in normalized.columns:
        normalized = normalized.rename(columns={"Class": "ISO Class"})
    if "Size" in normalized.columns:
        sizes = normalized["Size"].fillna("").astype(str).str.extract(
            r"(?P<width>\d+(?:\.\d+)?)\s*x\s*(?P<height>\d+(?:\.\d+)?)",
            expand=True,
        )
        if "Width (mm)" not in normalized.columns:
            normalized["Width (mm)"] = sizes["width"]
        if "Height (mm)" not in normalized.columns:
            normalized["Height (mm)"] = sizes["height"]
    if "Mass Efficiency" in normalized.columns and "Mass Eff. %" not in normalized.columns:
        normalized = normalized.rename(columns={"Mass Efficiency": "Mass Eff. %"})
    if "Price/filter" in normalized.columns and "Price/filter (VND)" not in normalized.columns:
        normalized = normalized.rename(columns={"Price/filter": "Price/filter (VND)"})
    if "DHC (g)" in normalized.columns and "DHC/filter direct (g)" not in normalized.columns:
        normalized = normalized.rename(columns={"DHC (g)": "DHC/filter direct (g)"})
    if "Area (m2)" in normalized.columns and "Media area/filter (m2)" not in normalized.columns:
        normalized = normalized.rename(columns={"Area (m2)": "Media area/filter (m2)"})
    normalized = normalized.reindex(columns=list(DISPLAY_COLUMNS.values()), fill_value="")
    if {"Width (mm)", "Height (mm)", "Media area/filter (m2)"}.issubset(normalized.columns):
        calculated_area = normalized.apply(
            lambda row: calculate_media_area_m2(row["Width (mm)"], row["Height (mm)"]),
            axis=1,
        )
        current_area = normalized["Media area/filter (m2)"].apply(_to_float)
        normalized["Media area/filter (m2)"] = [
            calculated if calculated > 0 and existing <= 0 else existing
            for calculated, existing in zip(calculated_area, current_area, strict=False)
        ]
    normalized = recalculate_mass_efficiency_dataframe(normalized)
    return normalized


def recalculate_media_area_dataframe(dataframe: pd.DataFrame) -> pd.DataFrame:
    normalized = normalize_display_dataframe(dataframe)
    if {"Width (mm)", "Height (mm)", "Media area/filter (m2)"}.issubset(normalized.columns):
        normalized["Media area/filter (m2)"] = normalized.apply(
            lambda row: calculate_media_area_m2(row["Width (mm)"], row["Height (mm)"]),
            axis=1,
        )
    normalized = recalculate_mass_efficiency_dataframe(normalized)
    return normalized


def sync_filter_database_editor(editor_key: str) -> None:
    editor_state = st.session_state.get(editor_key, {})
    current = normalize_display_dataframe(st.session_state.get("filter_database_data"))
    if not isinstance(editor_state, dict):
        return

    for row_index, changes in editor_state.get("edited_rows", {}).items():
        index = int(row_index)
        if index >= len(current):
            continue
        for column_name, value in changes.items():
            if column_name in current.columns:
                current.at[index, column_name] = value

    added_rows = editor_state.get("added_rows", [])
    if added_rows:
        current = pd.concat(
            [current, normalize_display_dataframe(pd.DataFrame(added_rows))],
            ignore_index=True,
        )

    deleted_rows = sorted(
        [int(index) for index in editor_state.get("deleted_rows", [])],
        reverse=True,
    )
    for index in deleted_rows:
        if index < len(current):
            current = current.drop(current.index[index])

    st.session_state.filter_database_data = recalculate_media_area_dataframe(
        current.reset_index(drop=True)
    )


def upsert_filter_database_row(row: dict[str, Any]) -> None:
    current = normalize_display_dataframe(st.session_state.get("filter_database_data"))
    filter_id = str(row.get("Filter ID", "")).strip()
    if not filter_id:
        st.warning("Filter ID is required before adding a filter record.")
        return

    new_row = normalize_display_dataframe(pd.DataFrame([row]))
    existing_matches = current["Filter ID"].astype(str).str.strip() == filter_id
    if existing_matches.any():
        current.loc[existing_matches, list(DISPLAY_COLUMNS.values())] = new_row.iloc[0].values
        st.success(f"Updated filter {filter_id}.")
    else:
        current = pd.concat([current, new_row], ignore_index=True)
        st.success(f"Added filter {filter_id}.")
    st.session_state.filter_database_data = normalize_display_dataframe(current)
    persist_filter_database_dataframe(current)
    st.session_state.filter_database_nonce = st.session_state.get(
        "filter_database_nonce", 0
    ) + 1


def _row_value(row: pd.Series | None, column: str, default: Any = "") -> Any:
    if row is None:
        return default
    value = row.get(column, default)
    if pd.isna(value):
        return default
    return value


def persist_filter_database_dataframe(dataframe: pd.DataFrame) -> None:
    st.session_state.filter_database_data = recalculate_media_area_dataframe(dataframe)
    st.session_state.filter_database_session_only = True


def remove_filter_database_rows(
    dataframe: pd.DataFrame,
    filter_ids: list[str],
) -> pd.DataFrame:
    normalized = normalize_display_dataframe(dataframe)
    ids_to_delete = {str(filter_id).strip() for filter_id in filter_ids if str(filter_id).strip()}
    if not ids_to_delete or "Filter ID" not in normalized.columns:
        return normalized
    keep_mask = ~normalized["Filter ID"].astype(str).str.strip().isin(ids_to_delete)
    return normalize_display_dataframe(normalized.loc[keep_mask].reset_index(drop=True))


def filter_record_ids(dataframe: pd.DataFrame) -> list[str]:
    normalized = normalize_display_dataframe(dataframe)
    return [
        str(value).strip()
        for value in normalized["Filter ID"].fillna("").tolist()
        if str(value).strip()
    ]


def _option_index(options: list[str], value: Any) -> int:
    text = str(value or "").strip()
    return options.index(text) if text in options else 0


def _missing_filter_record_fields(row: dict[str, Any]) -> list[str]:
    missing: list[str] = []
    required_text_fields = [
        "Filter ID",
        "Supplier",
        "Stage",
        "Filter model / description",
        "ISO Class",
    ]
    required_positive_fields = [
        "Qty/AHU",
        "Width (mm)",
        "Height (mm)",
        "Rated airflow/filter (m3/h)",
        "DHC/filter direct (g)",
        "Initial DP (Pa)",
        "Avg DP (Pa)",
        "Final DP (Pa)",
        "Price/filter (VND)",
    ]

    for field_name in required_text_fields:
        if not str(row.get(field_name, "")).strip():
            missing.append(field_name)

    for field_name in required_positive_fields:
        if _to_float(row.get(field_name, 0)) <= 0:
            missing.append(field_name)

    mass_efficiency, _source = calculate_effective_mass_efficiency_for_row(row)
    if mass_efficiency <= 0:
        missing.append("Mass Eff. % or ePM/Coarse")

    return missing


def render_quick_add_filter_form(current_df: pd.DataFrame) -> None:
    filter_ids = filter_record_ids(current_df)
    requested_edit_id = st.session_state.pop("filter_record_to_edit_requested", None)
    if requested_edit_id in filter_ids:
        st.session_state.filter_record_to_edit = requested_edit_id

    selected_existing = st.selectbox(
        "Edit existing filter",
        ["Create new filter"] + filter_ids,
        help="Choose an existing Filter ID to update it, or create a new filter.",
        key="filter_record_to_edit",
    )
    selected_row = None
    if selected_existing != "Create new filter":
        matches = current_df[current_df["Filter ID"].astype(str).str.strip() == selected_existing]
        if not matches.empty:
            selected_row = matches.iloc[0]
    form_key = selected_existing.replace(" ", "_").replace("/", "_").replace("\\", "_")

    with st.expander("Add or update one filter record", expanded=True):
        st.caption("Complete all required fields before saving. Notes are optional.")
        col1, col2, col3 = st.columns(3)
        filter_id = col1.text_input(
            "Filter ID",
            value=str(_row_value(selected_row, "Filter ID")),
            key=f"filter_id_{form_key}",
        )
        supplier = col2.text_input(
            "Supplier",
            value=str(_row_value(selected_row, "Supplier")),
            key=f"supplier_{form_key}",
        )
        model = col3.text_input(
            "Filter model / description",
            value=str(_row_value(selected_row, "Filter model / description")),
            key=f"model_{form_key}",
        )

        col4, col5, col6 = st.columns(3)
        stage = col4.selectbox(
            "Stage",
            STAGE_OPTIONS,
            index=_option_index(STAGE_OPTIONS, _row_value(selected_row, "Stage")),
            key=f"stage_{form_key}",
        )
        iso_class_options = _choice_options(ISO_CLASS_OPTIONS, current_df, "ISO Class")
        iso_class = col5.selectbox(
            "ISO Class",
            iso_class_options,
            index=_option_index(iso_class_options, _row_value(selected_row, "ISO Class")),
            key=f"iso_class_{form_key}",
            help="Reference only. This field helps identify the selected filter and is not used directly in LCC calculation.",
        )
        rated_airflow = col6.number_input(
            "Rated airflow/filter (m3/h)",
            min_value=0.0,
            value=float(_row_value(selected_row, "Rated airflow/filter (m3/h)", 0.0)),
            step=100.0,
            key=f"rated_airflow_{form_key}",
            help="Catalog or supplier rated airflow for one filter. Used for project airflow checking.",
        )
        parsed_iso = parse_iso_class_efficiencies(iso_class)

        col7, col8, col9 = st.columns(3)
        width_mm = col7.number_input(
            "Width (mm)",
            min_value=0.0,
            value=float(_row_value(selected_row, "Width (mm)", 0.0)),
            step=1.0,
            key=f"width_mm_{form_key}",
        )
        height_mm = col8.number_input(
            "Height (mm)",
            min_value=0.0,
            value=float(_row_value(selected_row, "Height (mm)", 0.0)),
            step=1.0,
            key=f"height_mm_{form_key}",
        )
        media_area = calculate_media_area_m2(width_mm, height_mm)
        col9.number_input(
            "Media area/filter (m2)",
            min_value=0.0,
            value=media_area,
            step=0.01,
            disabled=True,
            key=f"media_area_{form_key}",
            help="Automatically calculated from Width x Height / 1,000,000.",
        )

        col10, col11, col12 = st.columns(3)
        qty_per_ahu = col10.number_input(
            "Qty/AHU",
            min_value=0.0,
            value=float(_row_value(selected_row, "Qty/AHU", 0.0)),
            step=1.0,
            key=f"qty_per_ahu_{form_key}",
        )
        dhc_g = col11.number_input(
            "DHC/filter direct (g)",
            min_value=0.0,
            value=float(_row_value(selected_row, "DHC/filter direct (g)", 0.0)),
            step=50.0,
            key=f"dhc_g_{form_key}",
        )
        direct_mass_efficiency = col12.number_input(
            "Mass Eff. %",
            min_value=0.0,
            value=float(_row_value(selected_row, "Mass Eff. %", 0.0)),
            step=0.01,
            key=f"mass_efficiency_{form_key}",
            help="Enter measured mass efficiency if available. If blank/zero, the app estimates it from ePM/Coarse below.",
        )

        col13, col14, col15 = st.columns(3)
        initial_dp = col13.number_input(
            "Initial DP (Pa)",
            min_value=0.0,
            value=float(_row_value(selected_row, "Initial DP (Pa)", 0.0)),
            step=5.0,
            key=f"initial_dp_{form_key}",
        )
        avg_dp = col14.number_input(
            "Avg DP (Pa)",
            min_value=0.0,
            value=float(_row_value(selected_row, "Avg DP (Pa)", 0.0)),
            step=5.0,
            key=f"avg_dp_{form_key}",
        )
        final_dp = col15.number_input(
            "Final DP (Pa)",
            min_value=0.0,
            value=float(_row_value(selected_row, "Final DP (Pa)", 0.0)),
            step=5.0,
            key=f"final_dp_{form_key}",
        )

        col16, col17, col18, col19 = st.columns(4)
        epm1_percent = col16.number_input(
            "ePM1 %",
            min_value=0.0,
            max_value=100.0,
            value=float(_row_value(selected_row, "ePM1 %", 0.0) or parsed_iso.get("ePM1 %", 0.0)),
            step=1.0,
            key=f"epm1_{form_key}",
        )
        epm25_percent = col17.number_input(
            "ePM2.5 %",
            min_value=0.0,
            max_value=100.0,
            value=float(_row_value(selected_row, "ePM2.5 %", 0.0) or parsed_iso.get("ePM2.5 %", 0.0)),
            step=1.0,
            key=f"epm25_{form_key}",
        )
        epm10_percent = col18.number_input(
            "ePM10 %",
            min_value=0.0,
            max_value=100.0,
            value=float(_row_value(selected_row, "ePM10 %", 0.0) or parsed_iso.get("ePM10 %", 0.0)),
            step=1.0,
            key=f"epm10_{form_key}",
        )
        coarse_percent = col19.number_input(
            "ISO Coarse %",
            min_value=0.0,
            max_value=100.0,
            value=float(_row_value(selected_row, "ISO Coarse %", 0.0) or parsed_iso.get("ISO Coarse %", 0.0)),
            step=1.0,
            key=f"coarse_{form_key}",
        )

        mass_efficiency, mass_efficiency_source = effective_mass_efficiency(
            direct_mass_efficiency,
            epm1_percent,
            epm25_percent,
            epm10_percent,
            coarse_percent,
            current_outdoor_environment(),
            iso_class,
        )
        profile = environment_profile(current_outdoor_environment())
        st.caption(
            f"Effective Mass Eff.: {mass_efficiency * 100:,.2f}% ({mass_efficiency_source}) using {profile.label} dust profile."
        )

        col20, col21 = st.columns([1, 2])
        price = col20.number_input(
            "Price/filter (VND)",
            min_value=0.0,
            value=float(_row_value(selected_row, "Price/filter (VND)", 0.0)),
            step=10000.0,
            key=f"price_{form_key}",
        )
        notes = col21.text_input(
            "Notes",
            value=str(_row_value(selected_row, "Notes")),
            key=f"notes_{form_key}",
        )

        filter_record = {
            "Filter ID": filter_id,
            "Supplier": supplier,
            "Stage": stage,
            "Filter model / description": model,
            "ISO Class": iso_class,
            "Qty/AHU": qty_per_ahu,
            "Rated airflow/filter (m3/h)": rated_airflow,
            "Width (mm)": width_mm,
            "Height (mm)": height_mm,
            "Media area/filter (m2)": media_area,
            "DHC/filter direct (g)": dhc_g,
            "Initial DP (Pa)": initial_dp,
            "Avg DP (Pa)": avg_dp,
            "Final DP (Pa)": final_dp,
            "ePM1 %": epm1_percent,
            "ePM2.5 %": epm25_percent,
            "ePM10 %": epm10_percent,
            "ISO Coarse %": coarse_percent,
            "Mass Eff. %": mass_efficiency * 100.0,
            "Mass Eff. Source": mass_efficiency_source,
            "Price/filter (VND)": price,
            "Notes": notes,
        }
        missing_fields = _missing_filter_record_fields(filter_record)
        if missing_fields:
            st.info("Please complete: " + ", ".join(missing_fields))

        submitted = st.button(
            "Save filter record",
            width="stretch",
            disabled=bool(missing_fields),
        )
        if submitted:
            upsert_filter_database_row(filter_record)
            st.rerun()

    if selected_existing != "Create new filter":
        if st.button("Delete selected filter record", width="stretch"):
            remaining = current_df[
                current_df["Filter ID"].astype(str).str.strip() != selected_existing
            ].reset_index(drop=True)
            st.session_state.filter_database_data = normalize_display_dataframe(remaining)
            persist_filter_database_dataframe(remaining)
            st.session_state.filter_database_nonce = st.session_state.get(
                "filter_database_nonce", 0
            ) + 1
            st.success(f"Deleted filter {selected_existing}.")
            st.rerun()


def dataframe_to_records(dataframe: pd.DataFrame) -> list[FilterDatabaseRecord]:
    rename_map = {value: key for key, value in DISPLAY_COLUMNS.items()}
    internal = dataframe.rename(columns=rename_map)
    records: list[FilterDatabaseRecord] = []
    for _, row in internal.fillna("").iterrows():
        if not any(str(row.get(column, "")).strip() for column in DATABASE_COLUMNS):
            continue
        records.append(
            FilterDatabaseRecord(
                filter_id=str(row.get("filter_id", "")).strip(),
                supplier=str(row.get("supplier", "")).strip(),
                stage=str(row.get("stage", "")).strip(),
                model=str(row.get("model", "")).strip(),
                filter_class=str(row.get("filter_class", "")).strip(),
                qty_per_ahu=_to_float(row.get("qty_per_ahu", 1)),
                rated_airflow_m3_h_filter=_to_float(
                    row.get("rated_airflow_m3_h_filter", 0)
                ),
                width_mm=_to_float(row.get("width_mm", 0)),
                height_mm=_to_float(row.get("height_mm", 0)),
                media_area_m2=_to_float(row.get("media_area_m2", 0)),
                dhc_g=_to_float(row.get("dhc_g", 0)),
                initial_dp_pa=_to_float(row.get("initial_dp_pa", 0)),
                avg_dp_pa=_to_float(row.get("avg_dp_pa", 0)),
                final_dp_pa=_to_float(row.get("final_dp_pa", 0)),
                epm1_percent=_to_float(row.get("epm1_percent", 0)),
                epm25_percent=_to_float(row.get("epm25_percent", 0)),
                epm10_percent=_to_float(row.get("epm10_percent", 0)),
                coarse_percent=_to_float(row.get("coarse_percent", 0)),
                mass_efficiency=_to_float(row.get("mass_efficiency", 0)),
                mass_efficiency_source=str(row.get("mass_efficiency_source", "")).strip(),
                price_vnd_filter=_to_float(row.get("price_vnd_filter", 0)),
                notes=str(row.get("notes", "")).strip(),
            )
        )
    return records


def _to_float(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    try:
        return float(str(value).replace(",", "").replace("%", "").strip())
    except ValueError:
        return 0.0


def _normalize_column_name(name: object) -> str:
    return str(name or "").strip().lower().replace("_", " ")


def _known_column_count(columns: list[object]) -> int:
    known_aliases = {
        _normalize_column_name(alias)
        for aliases in COLUMN_ALIASES.values()
        for alias in aliases
    }
    return sum(1 for column in columns if _normalize_column_name(column) in known_aliases)


def _promote_detected_header_row(dataframe: pd.DataFrame) -> pd.DataFrame:
    if _known_column_count(list(dataframe.columns)) >= 3:
        return dataframe

    for row_index in range(min(len(dataframe), 12)):
        row_values = dataframe.iloc[row_index].tolist()
        normalized_values = {_normalize_column_name(value) for value in row_values}
        if "filter id" in normalized_values and "stage" in normalized_values:
            promoted = dataframe.iloc[row_index + 1 :].copy()
            promoted.columns = [str(value).strip() for value in row_values]
            return promoted.reset_index(drop=True)

    return dataframe


def normalize_uploaded_database(dataframe: pd.DataFrame) -> pd.DataFrame:
    source = _promote_detected_header_row(dataframe.dropna(how="all")).dropna(how="all").copy()
    source.columns = [str(column).strip() for column in source.columns]
    normalized = pd.DataFrame(columns=DATABASE_COLUMNS)

    source_lookup = {_normalize_column_name(column): column for column in source.columns}
    for target_column, aliases in COLUMN_ALIASES.items():
        matched_column = None
        for alias in aliases:
            if alias in source_lookup:
                matched_column = source_lookup[alias]
                break
        if matched_column is not None:
            normalized[target_column] = source[matched_column]
        else:
            normalized[target_column] = ""

    size_column = source_lookup.get("size") or source_lookup.get("dimension") or source_lookup.get("dimensions")
    if size_column is not None:
        sizes = source[size_column].fillna("").astype(str).str.extract(
            r"(?P<width>\d+(?:\.\d+)?)\s*x\s*(?P<height>\d+(?:\.\d+)?)",
            expand=True,
        )
        if not normalized["width_mm"].astype(str).str.strip().any():
            normalized["width_mm"] = sizes["width"]
        if not normalized["height_mm"].astype(str).str.strip().any():
            normalized["height_mm"] = sizes["height"]

    if "qty_per_ahu" in normalized.columns:
        blank_qty = normalized["qty_per_ahu"].astype(str).str.strip() == ""
        normalized.loc[blank_qty, "qty_per_ahu"] = 1

    for number_column in [
        "dhc_g",
        "qty_per_ahu",
        "rated_airflow_m3_h_filter",
        "width_mm",
        "height_mm",
        "media_area_m2",
        "initial_dp_pa",
        "avg_dp_pa",
        "final_dp_pa",
        "epm1_percent",
        "epm25_percent",
        "epm10_percent",
        "coarse_percent",
        "mass_efficiency",
        "price_vnd_filter",
    ]:
        normalized[number_column] = normalized[number_column].apply(_to_float)

    return recalculate_media_area_dataframe(normalized.rename(columns=DISPLAY_COLUMNS))


def read_excel_database(uploaded_file: Any, sheet_name: str) -> pd.DataFrame:
    raw = pd.read_excel(uploaded_file, sheet_name=sheet_name, header=None)
    return normalize_uploaded_database(raw)


def dataframe_to_excel_bytes(dataframe: pd.DataFrame) -> bytes:
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        normalize_display_dataframe(dataframe).to_excel(
            writer,
            index=False,
            sheet_name="Filter_Database",
        )
        list_sheet = writer.book.create_sheet("_Lists")
        for row_index, value in enumerate([option for option in STAGE_OPTIONS if option], start=2):
            list_sheet.cell(row=row_index, column=1, value=value)
        for row_index, value in enumerate([option for option in ISO_CLASS_OPTIONS if option], start=2):
            list_sheet.cell(row=row_index, column=2, value=value)
        list_sheet.sheet_state = "hidden"
        worksheet = writer.book["Filter_Database"]
        format_filter_database_worksheet(worksheet)
    return output.getvalue()


def format_filter_database_worksheet(worksheet: Any) -> None:
    header_fill = PatternFill("solid", fgColor="0F766E")
    required_fill = PatternFill("solid", fgColor="DCFCE7")
    dimension_fill = PatternFill("solid", fgColor="E0F2FE")
    pressure_fill = PatternFill("solid", fgColor="FEF9C3")
    cost_fill = PatternFill("solid", fgColor="FCE7F3")
    optional_fill = PatternFill("solid", fgColor="F8FAFC")
    warning_fill = PatternFill("solid", fgColor="FEE2E2")
    white_font = Font(color="FFFFFF", bold=True)
    body_font = Font(color="111827")
    thin_gray = Side(style="thin", color="CBD5E1")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)
    headers = [cell.value for cell in worksheet[1]]
    max_row = max(worksheet.max_row, 200)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(worksheet.max_column)}{max(worksheet.max_row, 2)}"
    worksheet.sheet_view.showGridLines = False

    column_widths = {
        "Filter ID": 18,
        "Supplier": 18,
        "Stage": 20,
        "Filter model / description": 34,
        "ISO Class": 18,
        "Qty/AHU": 12,
        "Rated airflow/filter (m3/h)": 24,
        "Width (mm)": 12,
        "Height (mm)": 12,
        "Media area/filter (m2)": 20,
        "DHC/filter direct (g)": 20,
        "Initial DP (Pa)": 15,
        "Avg DP (Pa)": 14,
        "Final DP (Pa)": 14,
        "ePM1 %": 12,
        "ePM2.5 %": 12,
        "ePM10 %": 12,
        "ISO Coarse %": 14,
        "Mass Eff. %": 13,
        "Mass Eff. Source": 22,
        "Price/filter (VND)": 18,
        "Notes": 28,
    }
    number_formats = {
        "Qty/AHU": "0",
        "Rated airflow/filter (m3/h)": "0",
        "Width (mm)": "0",
        "Height (mm)": "0",
        "Media area/filter (m2)": "0.00",
        "DHC/filter direct (g)": "0",
        "Initial DP (Pa)": "0",
        "Avg DP (Pa)": "0",
        "Final DP (Pa)": "0",
        "ePM1 %": "0.00",
        "ePM2.5 %": "0.00",
        "ePM10 %": "0.00",
        "ISO Coarse %": "0.00",
        "Mass Eff. %": "0.00",
        "Price/filter (VND)": '#,##0 "VND"',
    }
    column_groups = {
        "Filter ID": required_fill,
        "Supplier": required_fill,
        "Stage": required_fill,
        "Filter model / description": required_fill,
        "ISO Class": required_fill,
        "Qty/AHU": dimension_fill,
        "Rated airflow/filter (m3/h)": dimension_fill,
        "Width (mm)": dimension_fill,
        "Height (mm)": dimension_fill,
        "Media area/filter (m2)": dimension_fill,
        "DHC/filter direct (g)": pressure_fill,
        "Initial DP (Pa)": pressure_fill,
        "Avg DP (Pa)": pressure_fill,
        "Final DP (Pa)": pressure_fill,
        "ePM1 %": pressure_fill,
        "ePM2.5 %": pressure_fill,
        "ePM10 %": pressure_fill,
        "ISO Coarse %": pressure_fill,
        "Mass Eff. %": pressure_fill,
        "Mass Eff. Source": optional_fill,
        "Price/filter (VND)": cost_fill,
        "Notes": optional_fill,
    }
    comments = {
        "Avg DP (Pa)": "Recommended formula: round(((Initial DP + Final DP) / 2) * 1.1).",
        "Stage": "Choose a filter stage from the dropdown list.",
        "Mass Eff. %": "Used by LCC. If left blank in the app, it can be estimated from ePM/Coarse and Outdoor Environment.",
        "ePM1 %": "Optional ISO 16890 value used to estimate Mass Eff. when direct Mass Eff. is not available.",
        "ePM2.5 %": "Optional ISO 16890 value used to estimate Mass Eff. when direct Mass Eff. is not available.",
        "ePM10 %": "Optional ISO 16890 value used to estimate Mass Eff. when direct Mass Eff. is not available.",
        "ISO Coarse %": "Optional coarse arrestance value used to estimate Mass Eff. when direct Mass Eff. is not available.",
        "Price/filter (VND)": "Enter price for one filter, not total project price.",
        "Rated airflow/filter (m3/h)": "Catalog or supplier rated airflow for one filter. The app compares this with project airflow/filter.",
    }

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        if cell.value in comments:
            cell.comment = Comment(comments[cell.value], "LCC HVAC App")

    for column_index, header in enumerate(headers, start=1):
        letter = get_column_letter(column_index)
        worksheet.column_dimensions[letter].width = column_widths.get(str(header), 16)
        fill = column_groups.get(str(header), optional_fill)
        number_format = number_formats.get(str(header), "General")
        for row in range(2, max_row + 1):
            cell = worksheet[f"{letter}{row}"]
            cell.fill = fill
            cell.font = body_font
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.number_format = number_format
            cell.protection = Protection(locked=False)

    worksheet.row_dimensions[1].height = 34

    stage_column = headers.index("Stage") + 1 if "Stage" in headers else None
    if stage_column is not None:
        stage_letter = get_column_letter(stage_column)
        stage_count = len([option for option in STAGE_OPTIONS if option])
        validation = DataValidation(
            type="list",
            formula1=f"'_Lists'!$A$2:$A${stage_count + 1}",
            allow_blank=False,
            showErrorMessage=True,
        )
        validation.error = "Please choose one of the approved filter stages."
        validation.errorTitle = "Invalid stage"
        validation.prompt = "Select Pre-filter, Fine-filter, EPA / Final-filter, HEPA, or ULPA."
        validation.promptTitle = "Filter stage"
        worksheet.add_data_validation(validation)
        validation.add(f"{stage_letter}2:{stage_letter}{max_row}")

    iso_class_column = headers.index("ISO Class") + 1 if "ISO Class" in headers else None
    if iso_class_column is not None:
        iso_class_letter = get_column_letter(iso_class_column)
        iso_count = len([option for option in ISO_CLASS_OPTIONS if option])
        validation = DataValidation(
            type="list",
            formula1=f"'_Lists'!$B$2:$B${iso_count + 1}",
            allow_blank=False,
            showErrorMessage=True,
        )
        validation.error = "Please choose one of the approved ISO classes."
        validation.errorTitle = "Invalid ISO class"
        validation.prompt = "Select ISO Coarse, ISO ePM10, ISO ePM2.5, or ISO ePM1 class."
        validation.promptTitle = "ISO class"
        worksheet.add_data_validation(validation)
        validation.add(f"{iso_class_letter}2:{iso_class_letter}{max_row}")

    if all(header in headers for header in ["Width (mm)", "Height (mm)", "Media area/filter (m2)"]):
        width_letter = get_column_letter(headers.index("Width (mm)") + 1)
        height_letter = get_column_letter(headers.index("Height (mm)") + 1)
        media_area_letter = get_column_letter(headers.index("Media area/filter (m2)") + 1)
        for row in range(2, max_row + 1):
            worksheet[f"{media_area_letter}{row}"] = (
                f"=IF(AND({width_letter}{row}>0,{height_letter}{row}>0),"
                f"ROUND({width_letter}{row}*{height_letter}{row}/1000000,4),0)"
            )

    required_headers = ["Filter ID", "Supplier", "Stage", "Filter model / description", "ISO Class"]
    for required_header in required_headers:
        if required_header not in headers:
            continue
        letter = get_column_letter(headers.index(required_header) + 1)
        worksheet.conditional_formatting.add(
            f"{letter}2:{letter}{max_row}",
            FormulaRule(formula=[f'LEN(TRIM({letter}2))=0'], fill=warning_fill),
        )


def render_filter_record_delete_tools(saved_df: pd.DataFrame) -> pd.DataFrame:
    filter_ids = filter_record_ids(saved_df)
    if not filter_ids:
        return saved_df

    with st.expander("Delete filter records", expanded=False):
        selected_ids = st.multiselect(
            "Select Filter ID rows to delete",
            filter_ids,
            help="You can select one or many filter records, then delete them together.",
            key="filter_records_to_delete",
        )
        delete_col, count_col = st.columns([1, 2])
        delete_clicked = delete_col.button(
            "Delete selected rows",
            width="stretch",
            disabled=not selected_ids,
        )
        count_col.caption(f"{len(selected_ids)} row(s) selected.")
        if delete_clicked:
            remaining = remove_filter_database_rows(saved_df, selected_ids)
            st.session_state.filter_database_data = remaining
            persist_filter_database_dataframe(remaining)
            st.session_state.filter_database_nonce = st.session_state.get(
                "filter_database_nonce", 0
            ) + 1
            st.success(f"Deleted {len(selected_ids)} filter record(s).")
            st.rerun()
    return saved_df


def render_filter_record_edit_tools(saved_df: pd.DataFrame) -> None:
    filter_ids = filter_record_ids(saved_df)
    if not filter_ids:
        return

    with st.expander("Modify a row from Filter Records", expanded=False):
        selected_id = st.selectbox(
            "Select a Filter ID to modify",
            filter_ids,
            help="Choose a table row, then load it into the form above for editing.",
            key="filter_record_row_to_modify",
        )
        action_col, note_col = st.columns([1, 2])
        if action_col.button("Edit selected row", width="stretch"):
            st.session_state.filter_record_to_edit_requested = selected_id
            st.success(f"{selected_id} loaded into the edit form above.")
            st.rerun()
        note_col.caption("After editing, press Save filter record to update the row.")


def render_filter_database(records: list[FilterDatabaseRecord]) -> list[FilterDatabaseRecord]:
    st.subheader("Filter Database")
    st.caption(
        "Upload an Excel filter database or edit rows manually. Changes stay in your own session, so different users do not overwrite each other."
    )
    st.info(
        "Session database: upload or enter filters for this project, then download CSV/Excel if you want to keep a copy. "
        "The shared GitHub database is not changed by online users."
    )

    uploaded_file = st.file_uploader(
        "Upload Excel filter database",
        type=["xlsx", "xlsm", "xls"],
        help="If your workbook has a Filter_Database sheet, select it after upload.",
    )

    if uploaded_file is not None:
        try:
            excel_bytes = uploaded_file.getvalue()
            workbook = pd.ExcelFile(BytesIO(excel_bytes))
            default_sheet = (
                workbook.sheet_names.index("Filter_Database")
                if "Filter_Database" in workbook.sheet_names
                else 0
            )
            selected_sheet = st.selectbox(
                "Select sheet to import",
                workbook.sheet_names,
                index=default_sheet,
                key="filter_database_sheet",
            )
            if st.button("Import selected sheet", width="stretch"):
                imported = read_excel_database(BytesIO(excel_bytes), selected_sheet)
                st.session_state.filter_database_data = imported
                persist_filter_database_dataframe(imported)
                st.session_state.filter_database_nonce = st.session_state.get(
                    "filter_database_nonce", 0
                ) + 1
                st.success(f"Imported {len(imported)} rows from {selected_sheet}.")
                st.rerun()
        except Exception as exc:
            st.error(f"Cannot read uploaded Excel file: {exc}")

    current_df = st.session_state.get("filter_database_data")
    if current_df is None:
        current_df = records_to_dataframe(records)
        st.session_state.filter_database_data = normalize_display_dataframe(current_df)
    current_df = normalize_display_dataframe(st.session_state.filter_database_data)

    render_quick_add_filter_form(current_df)

    saved_df = normalize_display_dataframe(st.session_state.get("filter_database_data"))
    st.markdown("**Filter Records**")
    st.caption("Use the form above to add or edit filters. This table is for review and selection in Scenarios.")
    st.dataframe(saved_df, width="stretch", hide_index=True)
    render_filter_record_edit_tools(saved_df)
    saved_df = render_filter_record_delete_tools(saved_df)

    col1, col2 = st.columns([1, 2])
    if col1.button("Clear database", width="stretch"):
        empty_df = pd.DataFrame(columns=list(DISPLAY_COLUMNS.values()))
        st.session_state.filter_database_data = empty_df
        persist_filter_database_dataframe(empty_df)
        st.session_state.filter_database_nonce = st.session_state.get(
            "filter_database_nonce", 0
        ) + 1
        st.rerun()
    col2.download_button(
        "Download formatted Excel database",
        data=dataframe_to_excel_bytes(saved_df),
        file_name="filter_database.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        width="stretch",
    )

    return dataframe_to_records(saved_df)
